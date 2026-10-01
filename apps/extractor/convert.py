"""Conversion engines.

* sniff_kind: PDF / DOCX / image by magic bytes (anything else → Unsupported → 422)
* DOCX: python-docx, document order, tables (merged cells once, nested tables), headings → `#`
* PDF: pdfplumber text layer per page (ruled tables exactly, the rest clustered from word boxes); pages with fewer
  than MIN_TEXT_CHARS characters are rendered with pypdfium2 and OCR'd (ocr.py)
* image: OpenCV deskew → Tesseract
* Docling (layout model + TableFormer) is an optional extra: used for PDFs and images only when the package and
  its models are installed (DOCLING_ARTIFACTS_PATH), falling back to the basic engine on any error.
"""
from __future__ import annotations

import importlib.util
import logging
import os
import re
import time
import zipfile
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

import cv2
import docx
import numpy as np
import pdfplumber
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from PIL import Image

import ocr
from layout import Word, count_tables, rows_to_markdown, words_to_rows

log = logging.getLogger("extractor")
MIN_TEXT_CHARS = 50  # fewer characters on a PDF page ⇒ scanned page (cv-extraction skill)
MAX_PAGES = int(os.environ.get("EXTRACTOR_MAX_PAGES", "30"))


class Unsupported(Exception):
    """File type we do not handle, or too many pages → HTTP 422."""


@dataclass
class Doc:
    markdown: str
    pages: int
    text_layer: bool
    ocr_used: bool
    page_confidence: list[float]
    page1: np.ndarray | None = None  # BGR image of page 1 when we rendered/loaded one (for the face crop)
    pictures: list[Image.Image] = field(default_factory=list)  # embedded images the engine already extracted

    @property
    def tables(self) -> int:
        return count_tables(self.markdown)

    @property
    def confidence(self) -> float:
        return round(sum(self.page_confidence) / len(self.page_confidence), 3) if self.page_confidence else 0.0


def sniff_kind(path: Path) -> str:
    with path.open("rb") as f:
        head = f.read(8)
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        with zipfile.ZipFile(path) as z:
            if "word/document.xml" in z.namelist():
                return "docx"
        raise Unsupported(f"unsupported file type: {path.suffix or 'zip archive'}")
    try:
        with Image.open(path) as im:
            im.verify()
        return "image"
    except Exception:
        raise Unsupported(f"unsupported file type: {path.suffix or 'unknown'}") from None


def extract(kind: str, path: Path, langs: str) -> Doc:
    if kind == "docx":
        return extract_docx(path)
    if kind == "image":
        return extract_image(path, langs)
    return extract_pdf(path, langs)


# --- DOCX -------------------------------------------------------------------------------------------------------

def _table_rows(table: Table) -> list[list[str]]:
    rows: list[list[str]] = []
    for tr in table._tbl.tr_lst:  # raw rows: a merged cell is one w:tc, vertical continuations are skipped
        cells = [_Cell(tc, table) for tc in tr.tc_lst if tc.vMerge != "continue"]
        rows.append([c.text for c in cells])
        for c in cells:
            for nested in c.tables:
                rows.extend(_table_rows(nested))
    rows.append([])  # table boundary: keeps consecutive source tables apart in the Markdown
    return rows


def docx_pages(path: Path) -> int:
    with zipfile.ZipFile(path) as z:
        if "docProps/app.xml" in z.namelist():
            m = re.search(rb"<Pages>(\d+)</Pages>", z.read("docProps/app.xml"))
            if m:
                return max(1, int(m.group(1)))
    return 1


def extract_docx(path: Path) -> Doc:
    t = time.perf_counter()
    d = docx.Document(str(path))
    rows: list[list[str]] = []
    for el in d.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, d)
            text = p.text.strip()
            if not text:
                continue
            style = p.style.name if p.style is not None else ""
            if style.startswith("Heading") or style == "Title":
                text = "#" * (int(style[-1]) if style[-1].isdigit() else 1) + " " + text
            rows.append([text])
        elif el.tag == qn("w:tbl"):
            rows.extend(_table_rows(Table(el, d)))
    pages = docx_pages(path)
    log.info("docx: %d rows, %d page(s) in %.0f ms", len(rows), pages, (time.perf_counter() - t) * 1000)
    return Doc(rows_to_markdown(rows), pages, True, False, [1.0] * pages)


# --- PDF --------------------------------------------------------------------------------------------------------

def pdf_page_rows(page) -> tuple[list[list[str]], float]:
    """Rows of a text-layer page: ruled tables from pdfplumber, everything else clustered from word boxes.
    Confidence = share of characters that are real glyphs (not `(cid:n)` / U+FFFD placeholders)."""
    blocks: list[tuple[float, list[str]]] = []
    rest = page
    for t in page.find_tables():
        blocks += [(t.bbox[1] + i / 1000, [c or "" for c in row]) for i, row in enumerate(t.extract())]
        blocks.append((t.bbox[3] - 1e-6, []))  # table boundary
        rest = rest.outside_bbox(t.bbox)
    words = [Word(w["text"], w["x0"], w["x1"], w["top"], w["bottom"]) for w in rest.extract_words()]
    blocks += words_to_rows(words)
    blocks.sort(key=lambda b: b[0])
    bad = sum(c["text"].startswith("(cid:") or c["text"] == "�" for c in page.chars)
    return [cells for _, cells in blocks], round(1 - bad / max(1, len(page.chars)), 3)


def extract_pdf(path: Path, langs: str) -> Doc:
    with pdfplumber.open(str(path)) as pdf:
        n = len(pdf.pages)
        if n > MAX_PAGES:
            raise Unsupported(f"too many pages: {n} > {MAX_PAGES}")
        has_text = [len(p.chars) >= MIN_TEXT_CHARS for p in pdf.pages]
        if docling_enabled():
            try:
                return docling_extract("pdf", path, langs, text_layer=all(has_text))
            except Exception:
                log.exception("docling failed, falling back to the basic engine")
        rows: list[list[str]] = []
        confs: list[float] = []
        page1 = None
        for i, page in enumerate(pdf.pages):
            t = time.perf_counter()
            if has_text[i]:
                page_rows, conf = pdf_page_rows(page)
            else:
                img = ocr.render_pdf_page(path, i)
                page1 = img if i == 0 else page1
                words, conf = ocr.ocr_words(ocr.prepare(img, camera=False), langs)
                page_rows = [cells for _, cells in words_to_rows(words)]
            rows += page_rows
            confs.append(conf)
            log.info("page %d/%d: %s conf=%.2f in %.0f ms", i + 1, n, "text layer" if has_text[i] else f"ocr {langs}",
                     conf, (time.perf_counter() - t) * 1000)
    return Doc(rows_to_markdown(rows), n, any(has_text), not all(has_text), confs, page1)


# --- Photos of a page -------------------------------------------------------------------------------------------

def extract_image(path: Path, langs: str) -> Doc:
    t = time.perf_counter()
    img = ocr.deskew(ocr.load_image(path))
    if docling_enabled():
        flat = path.with_name(path.stem + ".deskewed.png")
        cv2.imwrite(str(flat), img)
        try:
            doc = docling_extract("image", flat, langs, text_layer=False)
            doc.page1 = img
            return doc
        except Exception:
            log.exception("docling failed, falling back to the basic engine")
    words, conf = ocr.ocr_words(ocr.prepare(img, camera=True), langs, adaptive=True)
    rows = [cells for _, cells in words_to_rows(words)]
    log.info("image: ocr %s conf=%.2f in %.0f ms", langs, conf, (time.perf_counter() - t) * 1000)
    return Doc(rows_to_markdown(rows), 1, False, True, [conf], page1=img)


# --- Optional Docling engine (models must be installed; see deploy/dockerfiles/extractor.Dockerfile) --------------

def docling_enabled() -> bool:
    mode = os.environ.get("EXTRACTOR_ENGINE", "auto")  # auto | basic | docling
    if mode == "basic" or importlib.util.find_spec("docling") is None:
        return False
    artifacts = os.environ.get("DOCLING_ARTIFACTS_PATH", "")
    return mode == "docling" or bool(artifacts and Path(artifacts).is_dir())


@lru_cache(maxsize=4)
def _docling_converter(do_ocr: bool, langs: str):
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import OcrMode, PdfPipelineOptions, TesseractCliOcrOptions
    from docling.document_converter import DocumentConverter, ImageFormatOption, PdfFormatOption

    ocr_options = TesseractCliOcrOptions(lang=langs.split("+"), **({"mode": OcrMode.FULL_PAGE} if do_ocr else {}))
    opts = PdfPipelineOptions(
        artifacts_path=os.environ.get("DOCLING_ARTIFACTS_PATH") or None,
        do_table_structure=True,
        generate_picture_images=True,
        images_scale=2.0,
        do_ocr=do_ocr,
        ocr_options=ocr_options,
    )
    return DocumentConverter(format_options={
        InputFormat.PDF: PdfFormatOption(pipeline_options=opts),
        InputFormat.IMAGE: ImageFormatOption(pipeline_options=opts),
    })


def docling_extract(kind: str, path: Path, langs: str, text_layer: bool) -> Doc:
    """Docling: layout model + TableFormer keep table structure even on scans; pictures come out as images."""
    t = time.perf_counter()
    do_ocr = not text_layer
    res = _docling_converter(do_ocr, langs).convert(str(path))
    d = res.document
    pages = d.num_pages() or 1
    confs = []
    for i in range(1, pages + 1):
        s = res.confidence.pages.get(i)
        v = (s.ocr_score if do_ocr else s.parse_score) if s is not None else float("nan")
        confs.append(round(float(v), 3) if v == v else 1.0)  # NaN = not scored
    pics = [p.get_image(d) for p in d.pictures]
    log.info("docling %s: %d page(s), ocr=%s in %.0f ms", kind, pages, do_ocr, (time.perf_counter() - t) * 1000)
    return Doc(d.export_to_markdown(), pages, text_layer, do_ocr, confs, pictures=[p for p in pics if p is not None])
