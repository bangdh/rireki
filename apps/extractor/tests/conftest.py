"""Synthetic 履歴書 documents generated at test time. No real people: the face is a public-domain portrait
(Grace Hopper, US Navy photo, from matplotlib's sample data), cropped 3:4 as an ID photo."""
from __future__ import annotations

import os
import shutil
from pathlib import Path

import boto3
import numpy as np
import pytest
from moto import mock_aws
from PIL import Image, ImageDraw, ImageFont

needs_tesseract = pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract binary not installed")
FIXTURES = Path(__file__).parent / "fixtures"
FACE = FIXTURES / "face.jpg"
JA_FONTS = [
    "/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf",
    "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
]
LATIN_FONTS = ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"]


def ja_font() -> str | None:
    return next((p for p in JA_FONTS if os.path.exists(p)), None)


# Rows mirror the Japanese render of the company template in app/candidate-detail.html.
HEAD = [["家族構成", "家族：4人（父・母・妹）", "状況", "就職活動中"]]
BASIC = [
    ["フリガナ", "グエン・バン・アン"],
    ["氏名", "NGUYEN VAN AN"],
    ["生年月日", "2002年3月15日", "24歳", "性別　男"],
    ["国籍", "ベトナム", "携帯電話番号", "+84 912 345 678"],
    ["メール", "an.nguyen@example.com"],
    ["フリガナ", "ベトナム　タインホア"],
    ["現住所", "ベトナム　タインホア省ホアンホア県"],
]
EDU = [
    ["年", "月", "学歴", "入学・卒業"],
    ["2017", "9", "ホアンホア第2高等学校（ベトナム）", "入学"],
    ["2020", "6", "ホアンホア第2高等学校（ベトナム）", "卒業"],
    ["2020", "9", "タインホア職業短期大学（ベトナム）", "入学"],
    ["2022", "6", "タインホア職業短期大学（ベトナム）", "卒業"],
]
WORK = [
    ["年", "月", "職歴", "入社・退職"],
    ["2022", "8", "ミンファット機械有限会社（ベトナム）", "入社"],
    ["2024", "5", "ミンファット機械有限会社（ベトナム）", "退職"],
    ["", "", "仕事内容：鉄骨フレームの溶接、品質チェック", ""],
    ["2024", "6", "サオベト研修センター（ベトナム）", "入社"],
    ["", "", "仕事内容：日本語研修、溶接実技訓練", ""],
]
CURRENT = [["現在", "現在 N3を勉強しています。"]]
LIC = [["年", "月", "免許・資格"], ["2023", "4", "溶接技能証明書 3G"], ["2025", "12", "JLPT N4"]]
SPOUSE = [["配偶者", "無", "配偶者の扶養義務", "無"]]
WISH = [["給与", "貴社規定に従います。"], ["勤務地", "全国どこでも大丈夫です。"], ["勤務時間", "会社スケジュールで大丈夫です。"]]
BODY = [["身長", "168 cm", "体重", "61 kg"], ["服のサイズ", "M", "肩（上半身）", "46 cm"], ["ウエスト（下半身）", "78 cm", "靴のサイズ", "26 cm"]]
NOTES = [
    ["宗教的に注意が必要な事項", "仏教です。特に注意が必要なことはありません。"],
    ["食べられないもの", "ありません。なんでも食べられます。"],
    ["アレルギー", "ありません。"],
    ["その他連絡事項", "—"],
]
HOBBIES = "趣味は料理とサッカーです。"
MOTIVATION = "専門学校で溶接を学び、2年間の実務でMIG/TIG溶接を担当してきました。"


def make_docx(path: Path, photo: bool = True) -> Path:
    import docx
    from docx.shared import Mm

    d = docx.Document()
    d.add_heading("履歴書", level=1)

    def table(rows):
        t = d.add_table(rows=0, cols=max(len(r) for r in rows))
        t.style = "Table Grid"
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = v
        return t

    table(HEAD)
    basic = table(BASIC)
    photo_cell = basic.cell(0, 3).merge(basic.cell(1, 3))  # vertically merged photo cell like the real template
    if photo:
        photo_cell.paragraphs[0].add_run().add_picture(str(FACE), width=Mm(30))
    d.add_heading("学歴", level=2)
    table(EDU)
    d.add_heading("職歴（アルバイト含む）", level=2)
    table(WORK)
    table(CURRENT)
    d.add_heading("免許・資格", level=2)
    table(LIC)
    d.add_heading("趣味・特技", level=2)
    d.add_paragraph(HOBBIES)
    d.add_heading("日本での就職志望動機、自己PR", level=2)
    d.add_paragraph(MOTIVATION)
    table(SPOUSE)
    table(WISH)
    table(BODY)
    table(NOTES)
    d.save(str(path))
    return path


def make_pdf(path: Path, photo: bool = True) -> Path:
    """Two-page digital PDF with ruled tables (reportlab + its built-in CJK font, so no font files are needed)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Image as RLImage
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = "HeiseiKakuGo-W5"
    pdfmetrics.registerFont(UnicodeCIDFont(font))
    cell = ParagraphStyle("cell", fontName=font, fontSize=9, leading=12)
    para = ParagraphStyle("para", fontName=font, fontSize=10, leading=14)
    head = ParagraphStyle("head", fontName=font, fontSize=13, leading=18, spaceBefore=8, spaceAfter=4)

    def table(rows, widths):
        n = len(widths)
        data = [[Paragraph(c, cell) for c in r + [""] * (n - len(r))] for r in rows]
        t = Table(data, colWidths=[w * mm for w in widths])
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return t

    story = [Paragraph("履歴書", head)]
    if photo:
        img = RLImage(str(FACE), width=30 * mm, height=40 * mm)
        img.hAlign = "RIGHT"
        story.append(img)
    story += [
        table(HEAD, [40, 60, 30, 40]), Spacer(1, 6), table(BASIC, [40, 60, 30, 40]),
        Paragraph("学歴", head), table(EDU, [20, 15, 100, 35]),
        Paragraph("職歴（アルバイト含む）", head), table(WORK, [20, 15, 100, 35]), table(CURRENT, [40, 130]),
        Paragraph("免許・資格", head), table(LIC, [20, 15, 135]),
        PageBreak(),
        Paragraph("趣味・特技", head), Paragraph(HOBBIES, para),
        Paragraph("日本での就職志望動機、自己PR", head), Paragraph(MOTIVATION, para),
        Spacer(1, 6), table(SPOUSE, [40, 45, 50, 35]), table(WISH, [45, 125]), table(BODY, [45, 40, 45, 40]), table(NOTES, [50, 120]),
    ]
    SimpleDocTemplate(str(path), pagesize=A4, title="履歴書").build(story)
    return path


SCAN_ROWS_JA = [
    ["フリガナ", "グエン・バン・アン"], ["氏名", "NGUYEN VAN AN"], ["生年月日", "2002年3月15日"], ["性別", "男"],
    ["国籍", "ベトナム"], ["携帯電話番号", "+84 912 345 678"], ["メール", "an.nguyen@example.com"],
    ["現住所", "ベトナム タインホア省"], ["配偶者", "無"], ["身長", "168 cm"], ["体重", "61 kg"],
    ["服のサイズ", "M"], ["靴のサイズ", "26 cm"], ["アレルギー", "ありません"],
]
SCAN_ROWS_EN = [["Name", "NGUYEN VAN AN"], ["Phone", "+84 912 345 678"], ["Email", "an.nguyen@example.com"], ["Height", "168 cm"]]


def _perspective_coeffs(src, dst):
    """PIL PERSPECTIVE coefficients mapping output points `src` to input points `dst`."""
    m = []
    for (x, y), (u, v) in zip(src, dst):
        m.append([x, y, 1, 0, 0, 0, -u * x, -u * y])
        m.append([0, 0, 0, x, y, 1, -v * x, -v * y])
    a = np.array(m, dtype=float)
    b = np.array(dst, dtype=float).reshape(8)
    return np.linalg.solve(a, b)


def make_scan(path: Path, photo: bool = True, warp: bool = False) -> tuple[Path, bool]:
    """A 'scanned' A4 page at 200 dpi: ruled form with label | value rows (Japanese when a CJK font is installed),
    the ID photo top-right. `warp` turns it into a phone photo: perspective distortion on a grey background.
    Returns (path, japanese)."""
    w, h = 1654, 2339
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    fp = ja_font()
    japanese = fp is not None
    font = ImageFont.truetype(fp or next(p for p in LATIN_FONTS if os.path.exists(p)), 34)
    d.text((120, 100), "履歴書" if japanese else "RIREKI CV", font=font, fill="black")
    y = 230
    for label, value in SCAN_ROWS_JA if japanese else SCAN_ROWS_EN:
        d.rectangle([110, y - 14, 1100, y + 54], outline="black", width=2)
        d.line([470, y - 14, 470, y + 54], fill="black", width=2)
        d.text((130, y), label, font=font, fill="black")
        d.text((490, y), value, font=font, fill="black")
        y += 90
    if photo:
        im.paste(Image.open(FACE).resize((330, 440)), (1200, 160))
    if warp:
        out = (w + 400, h + 400)
        quad = [(260, 230), (w + 120, 180), (w + 200, h + 220), (180, h + 300)]  # where the page corners land
        coeffs = _perspective_coeffs(quad, [(0, 0), (w, 0), (w, h), (0, h)])
        im = im.transform(out, Image.PERSPECTIVE, tuple(coeffs), Image.BICUBIC, fillcolor=(70, 70, 70))
    im.save(path)
    return path, japanese


def make_scan_pdf(path: Path, png: Path) -> Path:
    Image.open(png).convert("RGB").save(str(path), "PDF", resolution=200)  # image-only PDF: no text layer
    return path


@pytest.fixture
def s3(monkeypatch):
    """Mocked S3 (moto) with the originals bucket; the app's boto3 client must not point at a real endpoint."""
    for var in ("S3_ENDPOINT", "S3_ACCESS_KEY", "S3_SECRET_KEY", "S3_FORCE_PATH_STYLE", "S3_REGION", "AWS_PROFILE"):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setenv("AWS_ACCESS_KEY_ID", "testing")
    monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "testing")
    monkeypatch.setenv("AWS_DEFAULT_REGION", "us-east-1")
    with mock_aws():
        client = boto3.client("s3", region_name="us-east-1")
        client.create_bucket(Bucket="rireki-originals")
        yield client
